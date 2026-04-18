"""
Bot Utilities — Helper functions extracted from discord_memory_bot.py.

Contains file handling, crash reporting, lifecycle notifications, and
document detection utilities used by the Discord bot.
"""

from __future__ import annotations

import os
import sys
import traceback
from datetime import datetime
from pathlib import Path
from types import TracebackType
from typing import Any

import threading

import aiohttp

# Crash notification cooldown — prevents spamming Discord with identical errors
_crash_cooldown_lock = threading.Lock()
_crash_cooldown_until: float = 0.0  # time.time() when cooldown expires
_crash_suppressed_count: int = 0
CRASH_COOLDOWN_SECONDS: int = 300  # 5 minutes between crash notifications


def log(msg: str) -> None:
    """Print a timestamped log message."""
    ts = datetime.now().strftime("%H:%M:%S")
    print(f"[{ts}] {msg}", flush=True)


def build_crash_message(
    exc_type: type[BaseException], exc_value: BaseException, exc_tb: TracebackType | None
) -> str:
    """Build a detailed crash message with stack trace and local variables."""
    lines = []

    lines.append("**Stack Trace:**")
    lines.append("```python")
    tb_lines = traceback.format_exception(exc_type, exc_value, exc_tb)
    lines.extend([line.rstrip() for line in tb_lines])
    lines.append("```")

    lines.append("")
    lines.append("**Local Variables:**")
    lines.append("```python")

    tb = exc_tb
    while tb.tb_next:
        tb = tb.tb_next
    frame = tb.tb_frame

    for var_name, var_value in frame.f_locals.items():
        try:
            if isinstance(var_value, (type, type(sys))):
                continue
            value_str = repr(var_value)
            if len(value_str) > 100:
                value_str = value_str[:100] + "..."
            lines.append(f"{var_name} = {value_str}")
        except Exception:
            pass

    lines.append("```")

    result = "\n".join(lines)
    if len(result) > 1800:
        result = result[:1800] + "\n... [truncated]"
    return result


def write_crash_log(
    exc_type: type[BaseException],
    exc_value: BaseException,
    exc_tb: TracebackType | None,
    vault_path: str | Path,
) -> Path:
    """Write detailed crash information to the vault for later diagnosis.

    Includes the full stack trace and local variables from each frame.
    ``vault_path`` is the base Obsidian vault directory; the crash log is
    written to ``<vault>/LLM Memory/Permanent/crash_log.md``.
    """
    crash_file = Path(vault_path) / "LLM Memory" / "Permanent" / "crash_log.md"
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    lines = [
        "# Bot Crash Report",
        "",
        f"**Timestamp:** {timestamp}",
        f"**Exception Type:** {exc_type.__name__}",
        f"**Exception Message:** {exc_value}",
        "",
        "## Full Stack Trace",
        "```python",
    ]

    tb_lines = traceback.format_exception(exc_type, exc_value, exc_tb)
    lines.extend([line.rstrip() for line in tb_lines])
    lines.append("```")

    lines.append("")
    lines.append("## Local Variables by Frame")

    tb = exc_tb
    frame_num = 0
    while tb is not None:
        frame = tb.tb_frame
        lineno = tb.tb_lineno
        filename = frame.f_code.co_filename
        func_name = frame.f_code.co_name

        lines.append("")
        lines.append(f"### Frame {frame_num}: {func_name} ({filename}:{lineno})")
        lines.append("```python")

        for var_name, var_value in frame.f_locals.items():
            try:
                if isinstance(var_value, (type, type(sys))):
                    continue
                value_str = repr(var_value)
                if len(value_str) > 500:
                    value_str = value_str[:500] + "... [truncated]"
                lines.append(f"{var_name} = {value_str}")
            except Exception:
                lines.append(f"{var_name} = <unable to repr>")

        lines.append("```")
        tb = tb.tb_next
        frame_num += 1

    crash_file.parent.mkdir(parents=True, exist_ok=True)
    crash_file.write_text("\n".join(lines), encoding="utf-8")
    log(f"Crash log written to {crash_file}")

    return crash_file


def send_lifecycle_notification(event: str, details: str = "") -> None:
    """Send bot lifecycle event to Discord webhook.

    For crash events, enforces a cooldown to prevent spamming Discord
    when the same error repeats on every incoming message. The first
    crash is sent immediately; subsequent crashes within the cooldown
    window are suppressed and counted. When the cooldown expires, the
    next crash notification includes the suppressed count.
    """
    import time

    import requests

    global _crash_cooldown_until, _crash_suppressed_count

    webhook_url = os.environ.get("DISCORD_WEBHOOK_URL")
    if not webhook_url:
        return

    # Crash-specific cooldown to prevent channel spam
    if event == "crash":
        with _crash_cooldown_lock:
            now = time.time()
            if now < _crash_cooldown_until:
                _crash_suppressed_count += 1
                log(f"Crash notification suppressed ({_crash_suppressed_count} suppressed)")
                return
            # Cooldown expired or first crash — send it
            suppressed = _crash_suppressed_count
            _crash_suppressed_count = 0
            _crash_cooldown_until = now + CRASH_COOLDOWN_SECONDS

        if suppressed > 0:
            details = f"*({suppressed} additional crash(es) suppressed in the last {CRASH_COOLDOWN_SECONDS // 60} min)*\n{details}"

    icons = {
        "online": ":green_circle:",
        "offline": ":red_circle:",
        "crash": ":warning:",
        "restart": ":arrows_counterclockwise:",
    }
    icon = icons.get(event, ":robot:")

    message = f"{icon} **Bot {event.upper()}**"
    if details:
        message += f"\n{details}"
    message += f"\n*{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*"

    try:
        from .discord_rate_limit import retry_request

        # ``?wait=true`` makes Discord return the created message envelope so
        # we can record a clickable link for the /errors empty state. Only
        # worth the extra round-trip for ``online`` — crashes have their own
        # viewer on /errors and offline/restart are transient.
        post_url = webhook_url
        if event == "online" and "wait=" not in webhook_url:
            sep = "&" if "?" in webhook_url else "?"
            post_url = f"{webhook_url}{sep}wait=true"

        response = retry_request(requests.post, post_url, json={"content": message}, timeout=5)

        if event == "online":
            _record_online_notification(response)
    except Exception as e:
        log(f"Failed to send lifecycle notification: {e}")


def _record_online_notification(response: Any) -> None:
    """Parse the Discord webhook response and record the resulting message.

    When the webhook is posted with ``?wait=true`` Discord returns a JSON
    envelope containing ``id``, ``channel_id``, and (usually) a
    ``guild_id`` field from which a user-facing message URL can be built.
    Missing fields fall back to ``@me`` which is the legal Discord URL form
    for DMs / unattached webhooks and still renders a link in the UI.

    Any failure is swallowed — a missing URL is still a useful timestamp
    row, and the notification itself has already been sent successfully.
    """
    from .healthy_notifications import record_healthy_notification

    message_url: str | None = None
    try:
        if response is not None and getattr(response, "status_code", 500) < 300:
            body = response.json()
            message_id = body.get("id")
            channel_id = body.get("channel_id")
            guild_id = body.get("guild_id") or "@me"
            if message_id and channel_id:
                message_url = (
                    f"https://discord.com/channels/{guild_id}/{channel_id}/{message_id}"
                )
    except Exception as e:
        log(f"Could not parse webhook response for message URL: {e}")

    record_healthy_notification(message_url=message_url)


def detect_document_type(text: str) -> str | None:
    """Detect if pasted text is a resume/CV or other document worth storing.

    Only triggers for actual pasted documents (500+ chars), not short messages
    that merely mention the word 'resume'.
    """
    if len(text) < 500:
        return None

    text_lower = text.lower()

    resume_sections = [
        "experience", "education", "skills", "work history",
        "employment", "qualifications", "summary", "objective",
    ]
    section_count = sum(1 for s in resume_sections if s in text_lower)

    if any(word in text_lower[:200] for word in ["resume", "cv", "curriculum vitae"]):
        if section_count >= 1:
            return "resume"

    if section_count >= 3:
        return "resume"

    return None


async def extract_text_from_file(attachment: Any) -> str | None:
    """Download and extract text from Discord attachment."""
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(attachment.url) as resp:
                if resp.status != 200:
                    return None
                data = await resp.read()

        filename = attachment.filename.lower()

        if filename.endswith(".pdf"):
            import io
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text.strip() if text.strip() else None

        elif filename.endswith(".docx"):
            import io
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            return text.strip() if text.strip() else None

        elif filename.endswith((".txt", ".md", ".csv", ".log")):
            return data.decode("utf-8", errors="ignore").strip()

        elif filename.endswith((".py", ".js", ".ts", ".jsx", ".tsx", ".json", ".xml",
                                ".yaml", ".yml", ".html", ".css", ".sql", ".sh", ".bat",
                                ".toml", ".ini", ".cfg", ".env.example", ".gitignore",
                                ".java", ".go", ".rs", ".c", ".cpp", ".h", ".rb")):
            # Source code and config files — decode as UTF-8
            text = data.decode("utf-8", errors="ignore").strip()
            # Prefix with language hint for the LLM
            ext = filename.rsplit(".", 1)[-1] if "." in filename else ""
            return f"```{ext}\n{text}\n```" if text else None

        elif filename.endswith(".xlsx"):
            try:
                import io
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                lines = []
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    lines.append(f"## Sheet: {sheet}")
                    for row in ws.iter_rows(max_row=200, values_only=True):
                        cells = [str(c) if c is not None else "" for c in row]
                        lines.append(" | ".join(cells))
                wb.close()
                return "\n".join(lines).strip() if lines else None
            except ImportError:
                return None  # openpyxl not installed
            except Exception:
                return None

        else:
            return None

    except Exception as e:
        log(f"File extraction error: {e}")
        return None


# Image extensions supported by vision models
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp")


async def download_image(attachment: Any) -> bytes | None:
    """Download image bytes from Discord attachment."""
    try:
        async with aiohttp.ClientSession() as session:
            async with session.get(attachment.url) as resp:
                if resp.status != 200:
                    return None
                return await resp.read()
    except Exception as e:
        log(f"Image download error: {e}")
        return None


def is_image_attachment(attachment: Any) -> bool:
    """Check if attachment is an image we can process."""
    return attachment.filename.lower().endswith(IMAGE_EXTENSIONS)


# Extensions for file detection in responses
DEPLOYABLE_EXTENSIONS = {".html", ".htm"}
ATTACHABLE_EXTENSIONS = {".html", ".htm", ".txt", ".csv", ".json", ".md", ".pdf"}


def find_mentioned_files(response: str) -> list[Path]:
    """Scan a response for file paths that exist on disk.

    Looks for paths the bot mentions (e.g. "saved as japan_guide.html") and
    checks if they actually exist. Returns deduplicated list of real paths.
    """
    import re

    found: list[Path] = []
    seen: set[str] = set()

    for match in re.finditer(r"[A-Za-z]:\\[^\s\"'`<>|*?]+\.\w{2,5}", response):
        candidate = match.group(0).rstrip(".,;:)")
        _try_add_file(candidate, found, seen)

    for match in re.finditer(r"(?:\.?/[\w./-]+\.\w{2,5})", response):
        candidate = match.group(0).rstrip(".,;:)")
        _try_add_file(candidate, found, seen)

    for match in re.finditer(r"[\w][\w. -]*\.(?:html?|txt|csv|json|md|pdf)\b", response, re.IGNORECASE):
        candidate = match.group(0)
        for base in [
            Path.cwd(),
            Path(__file__).parent.parent,
            Path.home() / "Desktop",
            Path.home() / "OneDrive" / "Desktop",
            Path.home() / "Documents",
        ]:
            full = base / candidate
            if full.exists():
                _try_add_file(str(full), found, seen)
                break

    return found


def _try_add_file(path_str: str, found: list, seen: set) -> None:
    """Add a path to the file list if it's a real, attachable file."""
    try:
        if not path_str or len(path_str) > 500 or "\x00" in path_str:
            return
        p = Path(path_str).resolve()
        key = str(p).lower()
        if key in seen:
            return
        if p.exists() and p.is_file() and p.suffix.lower() in ATTACHABLE_EXTENSIONS:
            if p.stat().st_size < 10_000_000:
                found.append(p)
                seen.add(key)
    except Exception:
        pass
