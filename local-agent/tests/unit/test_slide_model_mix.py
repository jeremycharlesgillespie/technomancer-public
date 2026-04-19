"""Tests for scripts/slides/slide_06_model_mix.py (TK-627)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import agent.executor_runs_db as erdb


@pytest.fixture(autouse=True)
def _isolate_db(tmp_path, monkeypatch):
    """Point executor_runs_db at a temp dir and reset the connection cache."""
    db_path = tmp_path / "executor_runs.db"
    monkeypatch.setattr(erdb, "DB_DIR", tmp_path)
    monkeypatch.setattr(erdb, "DB_PATH", db_path)
    erdb._local.__dict__.pop("conn", None)
    yield
    conn = getattr(erdb._local, "conn", None)
    if conn:
        conn.close()
        erdb._local.__dict__.pop("conn", None)


def _insert_usage(rows: list[dict]) -> None:
    erdb.init_db()
    for row in rows:
        erdb.record_story_model_usage(
            story_key=row["story_key"],
            model=row["model"],
            call_count=row.get("call_count", 1),
            cost_usd=row.get("cost_usd", 0.0),
        )


def _chart_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART]


# ---------------------------------------------------------------------------
# Data slide tests
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self):
        _insert_usage([
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 10, "cost_usd": 0.05},
            {"story_key": "TK-1", "model": "claude-opus-4-6", "call_count": 5, "cost_usd": 0.80},
        ])
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        initial = len(prs.slides)
        add_slide(prs)
        assert len(prs.slides) == initial + 1

    def test_slide_has_title_text(self):
        _insert_usage([
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 10, "cost_usd": 0.05},
        ])
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame]
        assert any("Model" in t for t in texts)

    def test_slide_has_at_least_one_chart(self):
        _insert_usage([
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 10, "cost_usd": 0.05},
            {"story_key": "TK-2", "model": "claude-opus-4-6", "call_count": 5, "cost_usd": 0.80},
        ])
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(_chart_shapes(prs.slides[-1])) >= 1

    def test_headline_contains_worker_percentage(self):
        _insert_usage([
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 10, "cost_usd": 0.10},
            {"story_key": "TK-1", "model": "claude-opus-4-6", "call_count": 5, "cost_usd": 0.90},
        ])
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        assert "%" in all_text

    def test_saved_pptx_reopens_without_error(self, tmp_path):
        _insert_usage([
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 10, "cost_usd": 0.05},
            {"story_key": "TK-1", "model": "claude-opus-4-6", "call_count": 5, "cost_usd": 0.80},
        ])
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "model_mix.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_single_model_does_not_raise(self):
        _insert_usage([
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 8, "cost_usd": 0.20},
        ])
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1


# ---------------------------------------------------------------------------
# Empty / placeholder slide tests
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_adds_exactly_one_slide_when_empty(self):
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_does_not_raise(self):
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)

    def test_placeholder_contains_no_data_text(self):
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text

    def test_placeholder_has_title(self):
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        assert any("Model" in t for t in texts)

    def test_placeholder_pptx_reopens_without_error(self, tmp_path):
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "empty_model_mix.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_exception_in_db_falls_back_to_placeholder(self, monkeypatch):
        def _raise():
            raise RuntimeError("db down")

        monkeypatch.setattr(erdb, "get_all_story_model_usage", _raise)
        from scripts.slides.slide_06_model_mix import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text


# ---------------------------------------------------------------------------
# Registry tests
# ---------------------------------------------------------------------------


class TestRegistration:
    def test_slide_registered_at_slot_6(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_06_model_mix  # noqa: F401

        assert 6 in slide_registry

    def test_registered_value_is_add_slide(self):
        from scripts.slides import slide_registry
        from scripts.slides.slide_06_model_mix import add_slide

        assert slide_registry[6] is add_slide

    def test_registered_value_is_callable(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_06_model_mix  # noqa: F401

        assert callable(slide_registry[6])


# ---------------------------------------------------------------------------
# Internal helper: _aggregate_model_totals
# ---------------------------------------------------------------------------


class TestAggregateModelTotals:
    def test_sums_cost_per_model(self):
        from scripts.slides.slide_06_model_mix import _aggregate_model_totals

        rows = [
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 5, "cost_usd": 0.10},
            {"story_key": "TK-2", "model": "claude-haiku-4-5", "call_count": 3, "cost_usd": 0.05},
            {"story_key": "TK-1", "model": "claude-opus-4-6", "call_count": 2, "cost_usd": 0.80},
        ]
        result = _aggregate_model_totals(rows)
        haiku = next(r for r in result if "haiku" in r["model"])
        opus = next(r for r in result if "opus" in r["model"])
        assert abs(haiku["total_cost_usd"] - 0.15) < 1e-9
        assert abs(opus["total_cost_usd"] - 0.80) < 1e-9

    def test_counts_unique_stories(self):
        from scripts.slides.slide_06_model_mix import _aggregate_model_totals

        rows = [
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 5, "cost_usd": 0.10},
            {"story_key": "TK-2", "model": "claude-haiku-4-5", "call_count": 3, "cost_usd": 0.05},
        ]
        result = _aggregate_model_totals(rows)
        haiku = result[0]
        assert haiku["story_count"] == 2

    def test_computes_avg_calls_per_story(self):
        from scripts.slides.slide_06_model_mix import _aggregate_model_totals

        rows = [
            {"story_key": "TK-1", "model": "claude-opus-4-6", "call_count": 10, "cost_usd": 1.0},
            {"story_key": "TK-2", "model": "claude-opus-4-6", "call_count": 6, "cost_usd": 0.6},
        ]
        result = _aggregate_model_totals(rows)
        opus = result[0]
        # 2 stories, 16 total calls → 8.0 avg
        assert abs(opus["avg_calls_per_story"] - 8.0) < 1e-9

    def test_returns_sorted_by_cost_descending(self):
        from scripts.slides.slide_06_model_mix import _aggregate_model_totals

        rows = [
            {"story_key": "TK-1", "model": "claude-haiku-4-5", "call_count": 10, "cost_usd": 0.10},
            {"story_key": "TK-1", "model": "claude-opus-4-6", "call_count": 5, "cost_usd": 0.90},
        ]
        result = _aggregate_model_totals(rows)
        assert result[0]["total_cost_usd"] >= result[-1]["total_cost_usd"]

    def test_empty_rows_returns_empty(self):
        from scripts.slides.slide_06_model_mix import _aggregate_model_totals

        assert _aggregate_model_totals([]) == []


# ---------------------------------------------------------------------------
# Internal helper: _build_headline
# ---------------------------------------------------------------------------


class TestBuildHeadline:
    def test_contains_worker_percentage(self):
        from scripts.slides.slide_06_model_mix import _build_headline

        totals = [
            {"model": "claude-haiku-4-5", "total_cost_usd": 0.10},
            {"model": "claude-opus-4-6", "total_cost_usd": 0.90},
        ]
        result = _build_headline(totals)
        assert "90%" in result
        assert "worker" in result

    def test_zero_total_cost_returns_no_data_message(self):
        from scripts.slides.slide_06_model_mix import _build_headline

        totals = [{"model": "claude-haiku-4-5", "total_cost_usd": 0.0}]
        result = _build_headline(totals)
        assert "no cost data" in result.lower()

    def test_no_opus_shows_zero_percent(self):
        from scripts.slides.slide_06_model_mix import _build_headline

        totals = [{"model": "claude-haiku-4-5", "total_cost_usd": 1.0}]
        result = _build_headline(totals)
        assert "0%" in result


# ---------------------------------------------------------------------------
# Internal helper: _short_model_name
# ---------------------------------------------------------------------------


class TestShortModelName:
    def test_strips_claude_prefix(self):
        from scripts.slides.slide_06_model_mix import _short_model_name

        assert _short_model_name("claude-haiku-4-5-20251001") == "haiku-4-5"

    def test_keeps_three_segments(self):
        from scripts.slides.slide_06_model_mix import _short_model_name

        assert _short_model_name("claude-opus-4-6") == "opus-4-6"

    def test_no_prefix_passthrough(self):
        from scripts.slides.slide_06_model_mix import _short_model_name

        result = _short_model_name("opus-4-6")
        assert "opus" in result
