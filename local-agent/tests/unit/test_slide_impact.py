"""Tests for scripts/slides/slide_07_impact.py (TK-628)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import agent.daily_stats as ds
import agent.executor_runs_db as erdb


@pytest.fixture(autouse=True)
def _isolate_ds(tmp_path, monkeypatch):
    """Point daily_stats at a temp dir and reset the connection cache."""
    db_path = tmp_path / "daily_stats.db"
    monkeypatch.setattr(ds, "DB_DIR", tmp_path)
    monkeypatch.setattr(ds, "DB_PATH", db_path)
    ds._local.__dict__.pop("conn", None)
    yield
    conn = getattr(ds._local, "conn", None)
    if conn:
        conn.close()
        ds._local.__dict__.pop("conn", None)


@pytest.fixture(autouse=True)
def _isolate_erdb(tmp_path, monkeypatch):
    """Point executor_runs_db at a temp dir and reset the connection cache."""
    db_path = tmp_path / "executor_runs.db"
    monkeypatch.setattr(erdb, "DB_DIR", tmp_path)
    monkeypatch.setattr(erdb, "DB_PATH", db_path)
    erdb._local.__dict__.pop("conn", None)
    yield
    conn = getattr(erdb._local, "conn", None)
    if conn:
        conn.close()
        erdb._local.__dict__.pop("conn", None)


def _insert_ds_rows(rows: list[dict]) -> None:
    ds.init_db()
    conn = ds._get_conn()
    for row in rows:
        conn.execute(
            """INSERT OR REPLACE INTO daily_stats
               (date, project, shipped, loc_added, loc_removed)
               VALUES (?, ?, ?, ?, ?)""",
            (
                row["date"],
                row.get("project", "TK"),
                row.get("shipped", 1),
                row.get("loc_added", 0),
                row.get("loc_removed", 0),
            ),
        )
    conn.commit()


def _insert_run(
    jira_key: str,
    started_at: str,
    ended_at: str,
    status: str = "success",
) -> None:
    erdb.init_db()
    conn = erdb._get_conn()
    conn.execute(
        """INSERT INTO executor_runs
           (jira_key, branch, started_at, ended_at, status)
           VALUES (?, ?, ?, ?, ?)""",
        (jira_key, f"branch-{jira_key}", started_at, ended_at, status),
    )
    conn.commit()


def _chart_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART]


# ---------------------------------------------------------------------------
# Data slide tests
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self):
        _insert_ds_rows([
            {"date": "2026-04-01", "project": "TK", "loc_added": 120},
        ])
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        initial = len(prs.slides)
        add_slide(prs)
        assert len(prs.slides) == initial + 1

    def test_slide_has_title_text(self):
        _insert_ds_rows([
            {"date": "2026-04-01", "project": "TK", "loc_added": 50},
        ])
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame]
        assert any("Impact" in t or "LOC" in t for t in texts)

    def test_slide_has_at_least_one_chart(self):
        _insert_ds_rows([
            {"date": "2026-04-01", "project": "TK", "loc_added": 80},
            {"date": "2026-04-02", "project": "TK", "loc_added": 95},
        ])
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(_chart_shapes(prs.slides[-1])) >= 1

    def test_saved_pptx_reopens_without_error(self, tmp_path):
        _insert_ds_rows([
            {"date": "2026-04-01", "project": "TK", "loc_added": 100},
            {"date": "2026-04-02", "project": "TK", "loc_added": 150},
        ])
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "impact.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_peak_concurrent_shown_in_text(self):
        # Two overlapping runs → peak = 2
        _insert_run("TK-1", "2026-04-01T10:00:00", "2026-04-01T10:30:00")
        _insert_run("TK-2", "2026-04-01T10:10:00", "2026-04-01T10:40:00")
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        assert "concurrent" in all_text.lower()

    def test_multiple_projects_aggregated(self):
        _insert_ds_rows([
            {"date": "2026-04-01", "project": "TK", "loc_added": 100},
            {"date": "2026-04-01", "project": "FA", "loc_added": 50},
        ])
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1


# ---------------------------------------------------------------------------
# Empty / placeholder slide tests
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_adds_exactly_one_slide_when_empty(self):
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_does_not_raise(self):
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)

    def test_placeholder_contains_no_data_text(self):
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text

    def test_placeholder_has_title(self):
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        assert any("Impact" in t or "LOC" in t for t in texts)

    def test_placeholder_pptx_reopens_without_error(self, tmp_path):
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "empty_impact.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_exception_in_ds_falls_back_to_placeholder(self, monkeypatch):
        def _raise(*args, **kwargs):
            raise RuntimeError("db down")

        monkeypatch.setattr(ds, "get_rows", _raise)
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_exception_in_erdb_still_adds_slide(self, monkeypatch):
        _insert_ds_rows([{"date": "2026-04-01", "project": "TK", "loc_added": 50}])

        def _raise(*args, **kwargs):
            raise RuntimeError("db down")

        monkeypatch.setattr(erdb, "get_recent", _raise)
        from scripts.slides.slide_07_impact import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1


# ---------------------------------------------------------------------------
# Registry tests
# ---------------------------------------------------------------------------


class TestRegistration:
    def test_slide_registered_at_slot_7(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_07_impact  # noqa: F401

        assert 7 in slide_registry

    def test_registered_value_is_add_slide(self):
        from scripts.slides import slide_registry
        from scripts.slides.slide_07_impact import add_slide

        assert slide_registry[7] is add_slide

    def test_registered_value_is_callable(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_07_impact  # noqa: F401

        assert callable(slide_registry[7])


# ---------------------------------------------------------------------------
# Internal helper: _aggregate_loc
# ---------------------------------------------------------------------------


class TestAggregateLoc:
    def test_sums_across_projects(self):
        from scripts.slides.slide_07_impact import _aggregate_loc

        rows = [
            {"date": "2026-04-01", "project": "TK", "loc_added": 100},
            {"date": "2026-04-01", "project": "FA", "loc_added": 50},
            {"date": "2026-04-02", "project": "TK", "loc_added": 80},
        ]
        dates, per_day = _aggregate_loc(rows)
        assert per_day["2026-04-01"] == 150
        assert per_day["2026-04-02"] == 80

    def test_dates_sorted(self):
        from scripts.slides.slide_07_impact import _aggregate_loc

        rows = [
            {"date": "2026-04-03", "project": "TK", "loc_added": 10},
            {"date": "2026-04-01", "project": "TK", "loc_added": 20},
        ]
        dates, _ = _aggregate_loc(rows)
        assert dates == ["2026-04-01", "2026-04-03"]

    def test_empty_rows(self):
        from scripts.slides.slide_07_impact import _aggregate_loc

        dates, per_day = _aggregate_loc([])
        assert dates == []
        assert per_day == {}

    def test_missing_loc_treated_as_zero(self):
        from scripts.slides.slide_07_impact import _aggregate_loc

        rows = [{"date": "2026-04-01", "project": "TK"}]
        dates, per_day = _aggregate_loc(rows)
        assert per_day["2026-04-01"] == 0


# ---------------------------------------------------------------------------
# Internal helper: _running_totals
# ---------------------------------------------------------------------------


class TestRunningTotals:
    def test_basic_cumulative(self):
        from scripts.slides.slide_07_impact import _running_totals

        assert _running_totals([10, 20, 30]) == [10, 30, 60]

    def test_empty(self):
        from scripts.slides.slide_07_impact import _running_totals

        assert _running_totals([]) == []

    def test_single_value(self):
        from scripts.slides.slide_07_impact import _running_totals

        assert _running_totals([42]) == [42]


# ---------------------------------------------------------------------------
# Internal helper: _compute_peak_concurrent
# ---------------------------------------------------------------------------


class TestComputePeakConcurrent:
    def test_no_overlap(self):
        from scripts.slides.slide_07_impact import _compute_peak_concurrent

        rows = [
            {"started_at": "2026-04-01T10:00:00", "ended_at": "2026-04-01T10:15:00"},
            {"started_at": "2026-04-01T10:30:00", "ended_at": "2026-04-01T10:45:00"},
        ]
        assert _compute_peak_concurrent(rows) == 1

    def test_two_overlapping(self):
        from scripts.slides.slide_07_impact import _compute_peak_concurrent

        rows = [
            {"started_at": "2026-04-01T10:00:00", "ended_at": "2026-04-01T10:30:00"},
            {"started_at": "2026-04-01T10:10:00", "ended_at": "2026-04-01T10:40:00"},
        ]
        assert _compute_peak_concurrent(rows) == 2

    def test_three_overlapping(self):
        from scripts.slides.slide_07_impact import _compute_peak_concurrent

        rows = [
            {"started_at": "2026-04-01T10:00:00", "ended_at": "2026-04-01T11:00:00"},
            {"started_at": "2026-04-01T10:10:00", "ended_at": "2026-04-01T10:50:00"},
            {"started_at": "2026-04-01T10:20:00", "ended_at": "2026-04-01T10:40:00"},
        ]
        assert _compute_peak_concurrent(rows) == 3

    def test_empty_rows(self):
        from scripts.slides.slide_07_impact import _compute_peak_concurrent

        assert _compute_peak_concurrent([]) == 0

    def test_missing_timestamps_skipped(self):
        from scripts.slides.slide_07_impact import _compute_peak_concurrent

        rows = [
            {"started_at": None, "ended_at": "2026-04-01T10:30:00"},
            {"started_at": "2026-04-01T10:00:00", "ended_at": None},
            {"started_at": "2026-04-01T10:00:00", "ended_at": "2026-04-01T10:30:00"},
        ]
        assert _compute_peak_concurrent(rows) == 1
