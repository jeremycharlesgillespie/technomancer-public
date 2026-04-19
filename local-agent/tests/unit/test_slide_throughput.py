"""Tests for scripts/slides/slide_01_throughput.py (TK-622)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import agent.daily_stats as ds


@pytest.fixture(autouse=True)
def _isolate_db(tmp_path, monkeypatch):
    """Point daily_stats at a temp dir and reset the connection cache."""
    db_path = tmp_path / "daily_stats.db"
    monkeypatch.setattr(ds, "DB_DIR", tmp_path)
    monkeypatch.setattr(ds, "DB_PATH", db_path)
    ds._local.__dict__.pop("conn", None)
    yield
    conn = getattr(ds._local, "conn", None)
    if conn:
        conn.close()
        ds._local.__dict__.pop("conn", None)


def _insert_rows(rows: list[dict]) -> None:
    """Insert synthetic rows into the isolated daily_stats table."""
    ds.init_db()
    conn = ds._get_conn()
    for row in rows:
        conn.execute(
            "INSERT OR REPLACE INTO daily_stats (date, project, shipped) VALUES (?, ?, ?)",
            (row["date"], row["project"], row["shipped"]),
        )
    conn.commit()


def _chart_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART]


# ---------------------------------------------------------------------------
# Data slide tests
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self):
        _insert_rows([
            {"date": "2026-04-01", "project": "TK", "shipped": 3},
            {"date": "2026-04-02", "project": "TK", "shipped": 2},
            {"date": "2026-04-01", "project": "FA", "shipped": 1},
        ])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        initial = len(prs.slides)
        add_slide(prs)
        assert len(prs.slides) == initial + 1

    def test_slide_has_title_text(self):
        _insert_rows([{"date": "2026-04-01", "project": "TK", "shipped": 3}])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        slide = prs.slides[-1]
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        assert any("Stories Shipped" in t for t in texts)

    def test_slide_has_at_least_one_chart(self):
        _insert_rows([
            {"date": "2026-04-01", "project": "TK", "shipped": 3},
            {"date": "2026-04-02", "project": "TK", "shipped": 2},
        ])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(_chart_shapes(prs.slides[-1])) >= 1

    def test_slide_has_sparkline_chart(self):
        _insert_rows([
            {"date": "2026-04-01", "project": "TK", "shipped": 3},
            {"date": "2026-04-02", "project": "TK", "shipped": 2},
        ])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        # Expect bar chart + sparkline = 2 charts
        assert len(_chart_shapes(prs.slides[-1])) >= 2

    def test_headline_contains_total_shipped(self):
        _insert_rows([
            {"date": "2026-04-01", "project": "TK", "shipped": 5},
            {"date": "2026-04-02", "project": "TK", "shipped": 3},
        ])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        assert "8" in all_text  # 5 + 3 = 8

    def test_multi_project_produces_multiple_series(self):
        _insert_rows([
            {"date": "2026-04-01", "project": "TK", "shipped": 4},
            {"date": "2026-04-01", "project": "FA", "shipped": 2},
            {"date": "2026-04-02", "project": "TK", "shipped": 1},
        ])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        bar_chart = _chart_shapes(prs.slides[-1])[0].chart
        assert len(list(bar_chart.series)) >= 2

    def test_saved_pptx_reopens_without_error(self, tmp_path):
        """Written file can be re-opened (validates pptx structure)."""
        _insert_rows([{"date": "2026-04-01", "project": "TK", "shipped": 2}])
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "out.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1


# ---------------------------------------------------------------------------
# Empty / placeholder slide tests
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_adds_exactly_one_slide_when_empty(self):
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_does_not_raise(self):
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)  # must not raise

    def test_placeholder_contains_no_data_text(self):
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text

    def test_placeholder_pptx_reopens_without_error(self, tmp_path):
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "empty.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_exception_in_get_rows_falls_back_to_placeholder(self, monkeypatch):
        """If daily_stats.get_rows raises, we still get a placeholder slide."""
        import agent.daily_stats as _ds
        monkeypatch.setattr(_ds, "get_rows", lambda **kw: (_ for _ in ()).throw(RuntimeError("db down")))
        from scripts.slides.slide_01_throughput import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text


# ---------------------------------------------------------------------------
# Registry tests
# ---------------------------------------------------------------------------


class TestRegistration:
    def test_slide_registered_at_slot_1(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_01_throughput  # noqa: F401 — side effect registers it

        assert 1 in slide_registry

    def test_registered_value_is_add_slide(self):
        from scripts.slides import slide_registry
        from scripts.slides.slide_01_throughput import add_slide

        assert slide_registry[1] is add_slide

    def test_registered_value_is_callable(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_01_throughput  # noqa: F401

        assert callable(slide_registry[1])


# ---------------------------------------------------------------------------
# Internal helper tests
# ---------------------------------------------------------------------------


class TestAggregate:
    def test_aggregate_sums_shipped_per_project_per_date(self):
        from scripts.slides.slide_01_throughput import _aggregate

        rows = [
            {"date": "2026-04-01", "project": "TK", "shipped": 2},
            {"date": "2026-04-01", "project": "TK", "shipped": 3},
            {"date": "2026-04-02", "project": "FA", "shipped": 1},
        ]
        dates, per_project = _aggregate(rows)
        assert dates == ["2026-04-01", "2026-04-02"]
        assert per_project["TK"]["2026-04-01"] == 5
        assert per_project["FA"]["2026-04-02"] == 1

    def test_aggregate_returns_sorted_dates(self):
        from scripts.slides.slide_01_throughput import _aggregate

        rows = [
            {"date": "2026-04-03", "project": "TK", "shipped": 1},
            {"date": "2026-04-01", "project": "TK", "shipped": 1},
            {"date": "2026-04-02", "project": "TK", "shipped": 1},
        ]
        dates, _ = _aggregate(rows)
        assert dates == sorted(dates)


class TestCumulativeTotals:
    def test_running_sum_across_dates(self):
        from scripts.slides.slide_01_throughput import _cumulative_totals
        from collections import defaultdict

        dates = ["2026-04-01", "2026-04-02", "2026-04-03"]
        per_project = {
            "TK": defaultdict(int, {"2026-04-01": 2, "2026-04-02": 3, "2026-04-03": 1}),
        }
        result = _cumulative_totals(dates, per_project, ["TK"])
        assert result == [2, 5, 6]

    def test_multi_project_sums_across_projects(self):
        from scripts.slides.slide_01_throughput import _cumulative_totals
        from collections import defaultdict

        dates = ["2026-04-01", "2026-04-02"]
        per_project = {
            "TK": defaultdict(int, {"2026-04-01": 2, "2026-04-02": 1}),
            "FA": defaultdict(int, {"2026-04-01": 1, "2026-04-02": 2}),
        }
        result = _cumulative_totals(dates, per_project, ["TK", "FA"])
        assert result == [3, 6]
