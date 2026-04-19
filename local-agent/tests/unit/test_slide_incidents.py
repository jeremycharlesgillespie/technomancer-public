"""Tests for scripts/slides/slide_05_incidents.py."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from pptx import Presentation

import scripts.slides.slide_05_incidents as slide_module
from scripts.slides.slide_05_incidents import (
    _parse_incidents,
    _truncate,
    add_slide,
)

# ---------------------------------------------------------------------------
# Sample data
# ---------------------------------------------------------------------------

_SAMPLE_MD = """\
# Incidents Log

| Date | What broke | How detected | Fix commit |
|------|------------|--------------|------------|
| 2026-04-17 | splitter blocked recursive decomposition | AIM stalled on multi-step epics | TK-566 |
| 2026-04-18 | executor crashed on empty branch | test failure in CI | TK-600 |
| 2026-04-19 | rate limit hit on Jira API | 429 errors in logs | TK-610 |
"""

_EMPTY_MD = "# Incidents Log\n\nNo entries yet.\n"

_NO_TABLE_MD = "# Incidents Log\n\nSome prose, no table.\n"


def _make_prs() -> Presentation:
    return Presentation()


# ---------------------------------------------------------------------------
# _parse_incidents
# ---------------------------------------------------------------------------


class TestParseIncidents:
    def test_parses_all_rows(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        rows = _parse_incidents(f)
        assert len(rows) == 3

    def test_row_fields_correct(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        rows = _parse_incidents(f)
        assert rows[0] == (
            "2026-04-17",
            "splitter blocked recursive decomposition",
            "AIM stalled on multi-step epics",
            "TK-566",
        )

    def test_missing_file_returns_empty(self, tmp_path):
        rows = _parse_incidents(tmp_path / "nonexistent.md")
        assert rows == []

    def test_skips_header_row(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        rows = _parse_incidents(f)
        dates = [r[0] for r in rows]
        assert "Date" not in dates

    def test_skips_separator_row(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        rows = _parse_incidents(f)
        dates = [r[0] for r in rows]
        assert not any(d.startswith("---") for d in dates)

    def test_file_without_table_returns_empty(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_NO_TABLE_MD, encoding="utf-8")
        rows = _parse_incidents(f)
        assert rows == []


# ---------------------------------------------------------------------------
# _truncate
# ---------------------------------------------------------------------------


class TestTruncate:
    def test_short_string_unchanged(self):
        assert _truncate("hello", 10) == "hello"

    def test_exact_length_unchanged(self):
        assert _truncate("hello", 5) == "hello"

    def test_truncates_long_string(self):
        result = _truncate("a" * 100, 10)
        assert len(result) == 10

    def test_truncated_ends_with_ellipsis(self):
        result = _truncate("a" * 100, 10)
        assert result.endswith("\u2026")


# ---------------------------------------------------------------------------
# add_slide — with synthetic data
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        assert len(prs.slides) == 1

    def test_slide_has_title_textbox(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        slide = prs.slides[0]
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        assert any("Incidents" in t for t in texts)

    def test_slide_has_table(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        slide = prs.slides[0]
        tables = [s for s in slide.shapes if s.has_table]
        assert len(tables) == 1

    def test_table_has_four_columns(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
        assert len(tbl.columns) == 4

    def test_table_row_count_includes_header(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")  # 3 data rows
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
        assert len(tbl.rows) == 4  # header + 3 data rows

    def test_incidents_sorted_most_recent_first(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
        # First data row (row index 1) should hold the most recent date.
        first_date = tbl.cell(1, 0).text_frame.text
        assert first_date == "2026-04-19"

    def test_max_ten_incidents_shown(self, tmp_path):
        lines = ["# Incidents Log\n", "\n", "| Date | What broke | How detected | Fix commit |\n", "|------|------------|--------------|------------|\n"]
        for i in range(15):
            lines.append(f"| 2026-01-{i + 1:02d} | broke-{i} | detected-{i} | fix-{i} |\n")
        f = tmp_path / "incidents.md"
        f.write_text("".join(lines), encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
        assert len(tbl.rows) == 11  # header + 10 data rows


# ---------------------------------------------------------------------------
# add_slide — empty / missing data → placeholder
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_missing_file_adds_one_slide(self, tmp_path):
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", tmp_path / "missing.md"):
            add_slide(prs)
        assert len(prs.slides) == 1

    def test_empty_file_adds_one_slide(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_EMPTY_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_contains_no_data_text(self, tmp_path):
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", tmp_path / "missing.md"):
            add_slide(prs)
        slide = prs.slides[0]
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        assert any("No data yet" in t for t in texts)

    def test_placeholder_has_title(self, tmp_path):
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", tmp_path / "missing.md"):
            add_slide(prs)
        slide = prs.slides[0]
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        assert any("Incidents" in t for t in texts)

    def test_does_not_raise_on_missing_file(self, tmp_path):
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", tmp_path / "missing.md"):
            add_slide(prs)  # must not raise

    def test_does_not_raise_on_empty_file(self, tmp_path):
        f = tmp_path / "incidents.md"
        f.write_text(_EMPTY_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)  # must not raise


# ---------------------------------------------------------------------------
# pptx file validation
# ---------------------------------------------------------------------------


class TestPptxValidation:
    def test_slide_saves_and_reopens_with_data(self, tmp_path):
        """Saved pptx with data slide reopens cleanly (no 'content unreadable')."""
        f = tmp_path / "incidents.md"
        f.write_text(_SAMPLE_MD, encoding="utf-8")
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", f):
            add_slide(prs)
        out = tmp_path / "deck.pptx"
        prs.save(str(out))
        assert out.exists()
        prs2 = Presentation(str(out))
        assert len(prs2.slides) == 1

    def test_slide_saves_and_reopens_placeholder(self, tmp_path):
        """Saved pptx with placeholder slide reopens cleanly."""
        prs = _make_prs()
        with patch.object(slide_module, "INCIDENTS_PATH", tmp_path / "missing.md"):
            add_slide(prs)
        out = tmp_path / "deck_placeholder.pptx"
        prs.save(str(out))
        assert out.exists()
        prs2 = Presentation(str(out))
        assert len(prs2.slides) == 1
