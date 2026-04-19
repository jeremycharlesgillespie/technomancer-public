"""Tests for scripts.slides package skeleton (TK-620)."""

import importlib
import sys
from typing import Callable

import pytest
from pptx import Presentation


class TestSlideRegistry:
    """slide_registry is importable, empty by default, and typed correctly."""

    def setup_method(self):
        # Reload the module fresh so mutations from other tests don't leak.
        for mod_name in list(sys.modules):
            if mod_name.startswith("scripts.slides") or mod_name == "scripts.slides":
                del sys.modules[mod_name]

    def test_registry_is_dict(self):
        from scripts.slides import slide_registry

        assert isinstance(slide_registry, dict)

    def test_registry_is_empty_on_fresh_import(self):
        from scripts.slides import slide_registry

        assert len(slide_registry) == 0

    def test_registry_accepts_callable_values(self):
        from scripts.slides import slide_registry

        def dummy_builder(prs: Presentation) -> None:
            pass

        slide_registry[99] = dummy_builder
        assert 99 in slide_registry
        assert callable(slide_registry[99])

    def test_registry_key_is_int(self):
        from scripts.slides import slide_registry

        slide_registry[1] = lambda prs: None
        assert isinstance(list(slide_registry.keys())[0], int)


class TestCommonConstants:
    """Color, font, and geometry constants are present and correctly typed."""

    def test_color_constants_are_hex_strings(self):
        from scripts.slides.common import (
            COST_COLOR,
            FAILURE_COLOR,
            NEUTRAL_COLOR,
            SUCCESS_COLOR,
            THROUGHPUT_COLOR,
        )

        for name, value in [
            ("THROUGHPUT_COLOR", THROUGHPUT_COLOR),
            ("COST_COLOR", COST_COLOR),
            ("FAILURE_COLOR", FAILURE_COLOR),
            ("SUCCESS_COLOR", SUCCESS_COLOR),
            ("NEUTRAL_COLOR", NEUTRAL_COLOR),
        ]:
            assert isinstance(value, str), f"{name} should be a str"
            assert len(value) == 6, f"{name} should be 6 hex chars, got {value!r}"
            int(value, 16)  # raises ValueError if not valid hex

    def test_font_constants_are_strings(self):
        from scripts.slides.common import BODY_FONT, TITLE_FONT

        assert isinstance(TITLE_FONT, str) and TITLE_FONT
        assert isinstance(BODY_FONT, str) and BODY_FONT

    def test_size_constants_are_numeric(self):
        from scripts.slides.common import LABEL_SIZE_PT, SUBTITLE_SIZE_PT, TITLE_SIZE_PT

        for name, val in [
            ("TITLE_SIZE_PT", TITLE_SIZE_PT),
            ("SUBTITLE_SIZE_PT", SUBTITLE_SIZE_PT),
            ("LABEL_SIZE_PT", LABEL_SIZE_PT),
        ]:
            assert isinstance(val, (int, float)), f"{name} should be numeric"
            assert val > 0, f"{name} should be positive"

    def test_geometry_constants_are_positive_emu(self):
        from scripts.slides.common import (
            CHART_HEIGHT,
            CHART_LEFT,
            CHART_TOP,
            CHART_WIDTH,
            SLIDE_HEIGHT_EMU,
            SLIDE_WIDTH_EMU,
        )

        for name, val in [
            ("SLIDE_WIDTH_EMU", SLIDE_WIDTH_EMU),
            ("SLIDE_HEIGHT_EMU", SLIDE_HEIGHT_EMU),
            ("CHART_LEFT", CHART_LEFT),
            ("CHART_TOP", CHART_TOP),
            ("CHART_WIDTH", CHART_WIDTH),
            ("CHART_HEIGHT", CHART_HEIGHT),
        ]:
            assert isinstance(val, (int, float)), f"{name} should be numeric"
            assert val > 0, f"{name} should be positive"


class TestPptxInstalled:
    """python-pptx is importable and functional."""

    def test_pptx_importable(self):
        import pptx  # noqa: F401

    def test_can_create_presentation(self):
        prs = Presentation()
        assert prs is not None
