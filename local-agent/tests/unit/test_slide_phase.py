"""Tests for scripts/slides/slide_03_phase_breakdown.py (TK-624)."""

from __future__ import annotations

import json

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import agent.daily_stats as ds


@pytest.fixture(autouse=True)
def _isolate_db(tmp_path, monkeypatch):
    """Point daily_stats at a temp dir and reset the connection cache."""
    db_path = tmp_path / "daily_stats.db"
    monkeypatch.setattr(ds, "DB_DIR", tmp_path)
    monkeypatch.setattr(ds, "DB_PATH", db_path)
    ds._local.__dict__.pop("conn", None)
    yield
    conn = getattr(ds._local, "conn", None)
    if conn:
        conn.close()
        ds._local.__dict__.pop("conn", None)


def _insert_rows(rows: list[dict]) -> None:
    ds.init_db()
    conn = ds._get_conn()
    for row in rows:
        conn.execute(
            "INSERT OR REPLACE INTO daily_stats (date, project, shipped, cost_usd, phase_timings_json)"
            " VALUES (?, ?, ?, ?, ?)",
            (
                row["date"],
                row["project"],
                row.get("shipped", 0),
                row.get("cost_usd", 0.0),
                row.get("phase_timings_json"),
            ),
        )
    conn.commit()


def _chart_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART]


def _make_timings(**kwargs) -> str:
    """Build phase_timings_json with nested p50 values for named phases."""
    return json.dumps({phase: {"p50": secs} for phase, secs in kwargs.items()})


# ---------------------------------------------------------------------------
# Data slide tests
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 3,
             "phase_timings_json": _make_timings(claude_work=100.0, context_build=10.0)},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        initial = len(prs.slides)
        add_slide(prs)
        assert len(prs.slides) == initial + 1

    def test_slide_has_title_text(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=80.0)},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame]
        assert any("Phase" in t or "phase" in t for t in texts)

    def test_slide_has_chart(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=80.0, context_build=20.0)},
            {"date": "2026-04-16", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=90.0, context_build=10.0)},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(_chart_shapes(prs.slides[-1])) >= 1

    def test_headline_names_dominant_phase(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=200.0, context_build=5.0)},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        assert "claude_work" in all_text

    def test_null_phase_timings_does_not_raise(self):
        """Rows present but phase_timings_json is NULL — must not raise."""
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 2, "phase_timings_json": None},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_multi_project_aggregates_timings(self):
        timings = _make_timings(claude_work=100.0)
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "phase_timings_json": timings},
            {"date": "2026-04-15", "project": "FA", "phase_timings_json": timings},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_saved_pptx_reopens_without_error(self, tmp_path):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=80.0)},
        ])
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "phase.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1


# ---------------------------------------------------------------------------
# Empty / placeholder slide tests
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_adds_exactly_one_slide_when_empty(self):
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_does_not_raise(self):
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)

    def test_placeholder_contains_no_data_text(self):
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text

    def test_placeholder_has_title(self):
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        assert any("Phase" in t or "phase" in t for t in texts)

    def test_placeholder_pptx_reopens_without_error(self, tmp_path):
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "empty_phase.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_exception_in_get_rows_falls_back_to_placeholder(self, monkeypatch):
        monkeypatch.setattr(
            ds, "get_rows",
            lambda **kw: (_ for _ in ()).throw(RuntimeError("db down")),
        )
        from scripts.slides.slide_03_phase_breakdown import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text


# ---------------------------------------------------------------------------
# Registry tests
# ---------------------------------------------------------------------------


class TestRegistration:
    def test_slide_registered_at_slot_3(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_03_phase_breakdown  # noqa: F401

        assert 3 in slide_registry

    def test_registered_value_is_add_slide(self):
        from scripts.slides import slide_registry
        from scripts.slides.slide_03_phase_breakdown import add_slide

        assert slide_registry[3] is add_slide

    def test_registered_value_is_callable(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_03_phase_breakdown  # noqa: F401

        assert callable(slide_registry[3])


# ---------------------------------------------------------------------------
# Internal helper: _parse_phase_timings
# ---------------------------------------------------------------------------


class TestParsePhaseTimings:
    def test_parses_nested_p50_format(self):
        from scripts.slides.slide_03_phase_breakdown import _parse_phase_timings

        json_str = json.dumps({"claude_work": {"p50": 120.0, "p95": 300.0}})
        result = _parse_phase_timings(json_str)
        assert abs(result["claude_work"] - 120.0) < 1e-9

    def test_parses_flat_float_format(self):
        from scripts.slides.slide_03_phase_breakdown import _parse_phase_timings

        json_str = json.dumps({"claude_work": 120.0, "context_build": 10.0})
        result = _parse_phase_timings(json_str)
        assert abs(result["claude_work"] - 120.0) < 1e-9
        assert abs(result["context_build"] - 10.0) < 1e-9

    def test_returns_empty_on_invalid_json(self):
        from scripts.slides.slide_03_phase_breakdown import _parse_phase_timings

        assert _parse_phase_timings("not json") == {}

    def test_returns_empty_on_none(self):
        from scripts.slides.slide_03_phase_breakdown import _parse_phase_timings

        assert _parse_phase_timings(None) == {}

    def test_returns_empty_on_empty_string(self):
        from scripts.slides.slide_03_phase_breakdown import _parse_phase_timings

        assert _parse_phase_timings("") == {}


# ---------------------------------------------------------------------------
# Internal helper: _phase_percentages
# ---------------------------------------------------------------------------


class TestPhasePercentages:
    def test_percentages_sum_to_100(self):
        from scripts.slides.slide_03_phase_breakdown import _phase_percentages, PHASES

        rows = [
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=80.0, context_build=20.0)},
        ]
        dates, pcts = _phase_percentages(rows)
        total = sum(pcts[phase][0] for phase in PHASES)
        assert abs(total - 100.0) < 1e-6

    def test_empty_timings_returns_zeros(self):
        from scripts.slides.slide_03_phase_breakdown import _phase_percentages, PHASES

        rows = [{"date": "2026-04-15", "project": "TK", "phase_timings_json": None}]
        dates, pcts = _phase_percentages(rows)
        assert dates == ["2026-04-15"]
        for phase in PHASES:
            assert pcts[phase][0] == 0.0

    def test_returns_sorted_dates(self):
        from scripts.slides.slide_03_phase_breakdown import _phase_percentages

        rows = [
            {"date": "2026-04-17", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=100.0)},
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=100.0)},
        ]
        dates, _ = _phase_percentages(rows)
        assert dates == sorted(dates)

    def test_dominant_phase_matches_largest_input(self):
        from scripts.slides.slide_03_phase_breakdown import _phase_percentages, _dominant_phase

        rows = [
            {"date": "2026-04-15", "project": "TK",
             "phase_timings_json": _make_timings(claude_work=200.0, context_build=10.0)},
        ]
        dates, pcts = _phase_percentages(rows)
        dominant = _dominant_phase(dates, pcts)
        assert dominant == "claude_work"


# ---------------------------------------------------------------------------
# Internal helper: _dominant_phase
# ---------------------------------------------------------------------------


class TestDominantPhase:
    def test_returns_highest_avg_phase(self):
        from scripts.slides.slide_03_phase_breakdown import _dominant_phase, PHASES

        dates = ["2026-04-15"]
        pcts = {phase: [0.0] for phase in PHASES}
        pcts["context_build"] = [90.0]
        pcts["claude_work"] = [10.0]
        result = _dominant_phase(dates, pcts)
        assert result == "context_build"

    def test_returns_default_when_no_dates(self):
        from scripts.slides.slide_03_phase_breakdown import _dominant_phase, PHASES

        pcts = {phase: [] for phase in PHASES}
        result = _dominant_phase([], pcts)
        assert result == "claude_work"

    def test_averages_across_multiple_dates(self):
        from scripts.slides.slide_03_phase_breakdown import _dominant_phase, PHASES

        dates = ["2026-04-15", "2026-04-16"]
        pcts = {phase: [0.0, 0.0] for phase in PHASES}
        pcts["auto_commit"] = [60.0, 80.0]   # avg=70
        pcts["claude_work"] = [40.0, 20.0]   # avg=30
        result = _dominant_phase(dates, pcts)
        assert result == "auto_commit"
