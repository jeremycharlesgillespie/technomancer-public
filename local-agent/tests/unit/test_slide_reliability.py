"""Tests for scripts/slides/slide_04_reliability.py (TK-625)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import agent.daily_stats as ds


@pytest.fixture(autouse=True)
def _isolate_db(tmp_path, monkeypatch):
    """Point daily_stats at a temp dir and reset the connection cache."""
    db_path = tmp_path / "daily_stats.db"
    monkeypatch.setattr(ds, "DB_DIR", tmp_path)
    monkeypatch.setattr(ds, "DB_PATH", db_path)
    ds._local.__dict__.pop("conn", None)
    yield
    conn = getattr(ds._local, "conn", None)
    if conn:
        conn.close()
        ds._local.__dict__.pop("conn", None)


def _insert_rows(rows: list[dict]) -> None:
    ds.init_db()
    conn = ds._get_conn()
    for row in rows:
        conn.execute(
            "INSERT OR REPLACE INTO daily_stats "
            "(date, project, shipped, failed, first_attempt_success, "
            "splitter_child_success, splitter_child_fail) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                row["date"],
                row["project"],
                row.get("shipped", 0),
                row.get("failed", 0),
                row.get("first_attempt_success", 0),
                row.get("splitter_child_success"),
                row.get("splitter_child_fail"),
            ),
        )
    conn.commit()


def _chart_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART]


# ---------------------------------------------------------------------------
# Data slide tests
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 5,
             "first_attempt_success": 4, "splitter_child_success": 2, "splitter_child_fail": 1},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        initial = len(prs.slides)
        add_slide(prs)
        assert len(prs.slides) == initial + 1

    def test_slide_has_title_text(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 3,
             "first_attempt_success": 2, "splitter_child_success": 1, "splitter_child_fail": 0},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame]
        assert any("Reliability" in t or "reliability" in t for t in texts)

    def test_slide_has_two_charts(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 5,
             "first_attempt_success": 4, "splitter_child_success": 3, "splitter_child_fail": 1},
            {"date": "2026-04-16", "project": "TK", "shipped": 4,
             "first_attempt_success": 4, "splitter_child_success": 0, "splitter_child_fail": 0},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(_chart_shapes(prs.slides[-1])) == 2

    def test_headline_contains_recovery_percent(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 5,
             "first_attempt_success": 4, "splitter_child_success": 8, "splitter_child_fail": 2},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        # 8/(8+2)*100 = 80%
        assert "80" in all_text

    def test_null_splitter_data_does_not_raise(self):
        """Rows present but splitter columns are NULL — must not raise."""
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 3,
             "first_attempt_success": 2, "splitter_child_success": None, "splitter_child_fail": None},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_zero_shipped_does_not_divide_by_zero(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 0,
             "first_attempt_success": 0, "splitter_child_success": 0, "splitter_child_fail": 0},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_saved_pptx_reopens_without_error(self, tmp_path):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 5,
             "first_attempt_success": 4, "splitter_child_success": 3, "splitter_child_fail": 1},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "reliability.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_multi_project_aggregates_by_date(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 3,
             "first_attempt_success": 3, "splitter_child_success": 2, "splitter_child_fail": 0},
            {"date": "2026-04-15", "project": "FA", "shipped": 2,
             "first_attempt_success": 1, "splitter_child_success": 1, "splitter_child_fail": 1},
        ])
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1


# ---------------------------------------------------------------------------
# Empty / placeholder slide tests
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_adds_exactly_one_slide_when_empty(self):
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_does_not_raise(self):
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)

    def test_placeholder_contains_no_data_text(self):
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text

    def test_placeholder_has_title(self):
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        assert any("Reliability" in t or "reliability" in t for t in texts)

    def test_placeholder_pptx_reopens_without_error(self, tmp_path):
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "empty_reliability.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_exception_in_get_rows_falls_back_to_placeholder(self, monkeypatch):
        monkeypatch.setattr(
            ds, "get_rows",
            lambda **kw: (_ for _ in ()).throw(RuntimeError("db down")),
        )
        from scripts.slides.slide_04_reliability import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text


# ---------------------------------------------------------------------------
# Registry tests
# ---------------------------------------------------------------------------


class TestRegistration:
    def test_slide_registered_at_slot_4(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_04_reliability  # noqa: F401

        assert 4 in slide_registry

    def test_registered_value_is_add_slide(self):
        from scripts.slides import slide_registry
        from scripts.slides.slide_04_reliability import add_slide

        assert slide_registry[4] is add_slide

    def test_registered_value_is_callable(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_04_reliability  # noqa: F401

        assert callable(slide_registry[4])


# ---------------------------------------------------------------------------
# Internal helper: _first_attempt_rates
# ---------------------------------------------------------------------------


class TestFirstAttemptRates:
    def test_basic_rate_calculation(self):
        from scripts.slides.slide_04_reliability import _first_attempt_rates

        rows = [{"date": "2026-04-15", "project": "TK", "shipped": 4, "first_attempt_success": 3}]
        dates, rates = _first_attempt_rates(rows)
        assert dates == ["2026-04-15"]
        assert abs(rates[0] - 75.0) < 1e-6

    def test_zero_shipped_gives_zero_rate(self):
        from scripts.slides.slide_04_reliability import _first_attempt_rates

        rows = [{"date": "2026-04-15", "project": "TK", "shipped": 0, "first_attempt_success": 0}]
        _, rates = _first_attempt_rates(rows)
        assert rates[0] == 0.0

    def test_aggregates_multiple_projects_same_date(self):
        from scripts.slides.slide_04_reliability import _first_attempt_rates

        rows = [
            {"date": "2026-04-15", "project": "TK", "shipped": 4, "first_attempt_success": 4},
            {"date": "2026-04-15", "project": "FA", "shipped": 4, "first_attempt_success": 0},
        ]
        dates, rates = _first_attempt_rates(rows)
        assert len(dates) == 1
        assert abs(rates[0] - 50.0) < 1e-6

    def test_returns_sorted_dates(self):
        from scripts.slides.slide_04_reliability import _first_attempt_rates

        rows = [
            {"date": "2026-04-17", "project": "TK", "shipped": 2, "first_attempt_success": 2},
            {"date": "2026-04-15", "project": "TK", "shipped": 2, "first_attempt_success": 1},
        ]
        dates, _ = _first_attempt_rates(rows)
        assert dates == sorted(dates)

    def test_one_hundred_percent_rate(self):
        from scripts.slides.slide_04_reliability import _first_attempt_rates

        rows = [{"date": "2026-04-15", "project": "TK", "shipped": 5, "first_attempt_success": 5}]
        _, rates = _first_attempt_rates(rows)
        assert abs(rates[0] - 100.0) < 1e-6


# ---------------------------------------------------------------------------
# Internal helper: _splitter_children
# ---------------------------------------------------------------------------


class TestSplitterChildren:
    def test_basic_counts(self):
        from scripts.slides.slide_04_reliability import _splitter_children

        rows = [
            {"date": "2026-04-15", "project": "TK",
             "splitter_child_success": 5, "splitter_child_fail": 2},
        ]
        dates, shipped, failed = _splitter_children(rows)
        assert dates == ["2026-04-15"]
        assert shipped == [5]
        assert failed == [2]

    def test_null_values_treated_as_zero(self):
        from scripts.slides.slide_04_reliability import _splitter_children

        rows = [
            {"date": "2026-04-15", "project": "TK",
             "splitter_child_success": None, "splitter_child_fail": None},
        ]
        _, shipped, failed = _splitter_children(rows)
        assert shipped == [0]
        assert failed == [0]

    def test_aggregates_by_date(self):
        from scripts.slides.slide_04_reliability import _splitter_children

        rows = [
            {"date": "2026-04-15", "project": "TK",
             "splitter_child_success": 3, "splitter_child_fail": 1},
            {"date": "2026-04-15", "project": "FA",
             "splitter_child_success": 2, "splitter_child_fail": 2},
        ]
        dates, shipped, failed = _splitter_children(rows)
        assert len(dates) == 1
        assert shipped == [5]
        assert failed == [3]

    def test_returns_sorted_dates(self):
        from scripts.slides.slide_04_reliability import _splitter_children

        rows = [
            {"date": "2026-04-17", "project": "TK",
             "splitter_child_success": 1, "splitter_child_fail": 0},
            {"date": "2026-04-15", "project": "TK",
             "splitter_child_success": 2, "splitter_child_fail": 1},
        ]
        dates, _, _ = _splitter_children(rows)
        assert dates == sorted(dates)


# ---------------------------------------------------------------------------
# Internal helper: _overall_recovery_rate
# ---------------------------------------------------------------------------


class TestOverallRecoveryRate:
    def test_basic_rate(self):
        from scripts.slides.slide_04_reliability import _overall_recovery_rate

        rows = [{"splitter_child_success": 8, "splitter_child_fail": 2}]
        rate = _overall_recovery_rate(rows)
        assert abs(rate - 80.0) < 1e-6

    def test_zero_children_returns_zero(self):
        from scripts.slides.slide_04_reliability import _overall_recovery_rate

        rows = [{"splitter_child_success": 0, "splitter_child_fail": 0}]
        rate = _overall_recovery_rate(rows)
        assert rate == 0.0

    def test_null_children_returns_zero(self):
        from scripts.slides.slide_04_reliability import _overall_recovery_rate

        rows = [{"splitter_child_success": None, "splitter_child_fail": None}]
        rate = _overall_recovery_rate(rows)
        assert rate == 0.0

    def test_aggregates_across_multiple_rows(self):
        from scripts.slides.slide_04_reliability import _overall_recovery_rate

        rows = [
            {"splitter_child_success": 6, "splitter_child_fail": 4},
            {"splitter_child_success": 4, "splitter_child_fail": 6},
        ]
        rate = _overall_recovery_rate(rows)
        assert abs(rate - 50.0) < 1e-6

    def test_one_hundred_percent_recovery(self):
        from scripts.slides.slide_04_reliability import _overall_recovery_rate

        rows = [{"splitter_child_success": 10, "splitter_child_fail": 0}]
        rate = _overall_recovery_rate(rows)
        assert abs(rate - 100.0) < 1e-6
