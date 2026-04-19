"""Tests for scripts/slides/slide_02_cost.py (TK-623)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import agent.daily_stats as ds


@pytest.fixture(autouse=True)
def _isolate_db(tmp_path, monkeypatch):
    """Point daily_stats at a temp dir and reset the connection cache."""
    db_path = tmp_path / "daily_stats.db"
    monkeypatch.setattr(ds, "DB_DIR", tmp_path)
    monkeypatch.setattr(ds, "DB_PATH", db_path)
    ds._local.__dict__.pop("conn", None)
    yield
    conn = getattr(ds._local, "conn", None)
    if conn:
        conn.close()
        ds._local.__dict__.pop("conn", None)


def _insert_rows(rows: list[dict]) -> None:
    ds.init_db()
    conn = ds._get_conn()
    for row in rows:
        conn.execute(
            "INSERT OR REPLACE INTO daily_stats (date, project, shipped, cost_usd)"
            " VALUES (?, ?, ?, ?)",
            (row["date"], row["project"], row.get("shipped", 0), row.get("cost_usd", 0.0)),
        )
    conn.commit()


def _chart_shapes(slide):
    return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.CHART]


# ---------------------------------------------------------------------------
# Data slide tests
# ---------------------------------------------------------------------------


class TestAddSlideWithData:
    def test_adds_exactly_one_slide(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 3, "cost_usd": 0.90},
            {"date": "2026-04-17", "project": "TK", "shipped": 2, "cost_usd": 0.10},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        initial = len(prs.slides)
        add_slide(prs)
        assert len(prs.slides) == initial + 1

    def test_slide_has_title_text(self):
        _insert_rows([{"date": "2026-04-15", "project": "TK", "shipped": 2, "cost_usd": 0.50}])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame]
        assert any("Cost" in t for t in texts)

    def test_slide_has_line_chart(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 2, "cost_usd": 0.50},
            {"date": "2026-04-16", "project": "TK", "shipped": 3, "cost_usd": 0.30},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(_chart_shapes(prs.slides[-1])) >= 1

    def test_headline_shows_after_cost(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 2, "cost_usd": 1.00},
            {"date": "2026-04-17", "project": "TK", "shipped": 4, "cost_usd": 0.20},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        # After avg: 0.20 / 4 = 0.05 → "$0.0500"
        assert "0.05" in all_text

    def test_annotation_labels_appear_when_switch_date_in_data(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 2, "cost_usd": 1.00},
            {"date": "2026-04-17", "project": "TK", "shipped": 2, "cost_usd": 0.10},
            {"date": "2026-04-18", "project": "TK", "shipped": 2, "cost_usd": 0.08},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame
        )
        assert "Before" in all_text
        assert "After" in all_text

    def test_no_annotation_when_switch_date_absent(self):
        _insert_rows([
            {"date": "2026-04-10", "project": "TK", "shipped": 2, "cost_usd": 0.50},
            {"date": "2026-04-11", "project": "TK", "shipped": 2, "cost_usd": 0.40},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)  # must not raise

    def test_shipped_zero_does_not_raise(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 0, "cost_usd": 0.50},
            {"date": "2026-04-16", "project": "TK", "shipped": 2, "cost_usd": 0.30},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_multi_project_aggregates_to_single_series(self):
        _insert_rows([
            {"date": "2026-04-15", "project": "TK", "shipped": 2, "cost_usd": 0.40},
            {"date": "2026-04-15", "project": "FA", "shipped": 2, "cost_usd": 0.20},
        ])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        chart = _chart_shapes(prs.slides[-1])[0].chart
        assert len(list(chart.series)) == 1

    def test_saved_pptx_reopens_without_error(self, tmp_path):
        _insert_rows([{"date": "2026-04-15", "project": "TK", "shipped": 2, "cost_usd": 0.50}])
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "cost.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1


# ---------------------------------------------------------------------------
# Empty / placeholder slide tests
# ---------------------------------------------------------------------------


class TestAddSlideEmpty:
    def test_adds_exactly_one_slide_when_empty(self):
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1

    def test_placeholder_does_not_raise(self):
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)

    def test_placeholder_contains_no_data_text(self):
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text

    def test_placeholder_has_title(self):
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
        assert any("Cost" in t for t in texts)

    def test_placeholder_pptx_reopens_without_error(self, tmp_path):
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        path = tmp_path / "empty_cost.pptx"
        prs.save(str(path))
        reopened = Presentation(str(path))
        assert len(reopened.slides) == 1

    def test_exception_in_get_rows_falls_back_to_placeholder(self, monkeypatch):
        monkeypatch.setattr(
            ds, "get_rows",
            lambda **kw: (_ for _ in ()).throw(RuntimeError("db down")),
        )
        from scripts.slides.slide_02_cost import add_slide

        prs = Presentation()
        add_slide(prs)
        assert len(prs.slides) == 1
        all_text = " ".join(
            sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
        )
        assert "No data yet" in all_text


# ---------------------------------------------------------------------------
# Registry tests
# ---------------------------------------------------------------------------


class TestRegistration:
    def test_slide_registered_at_slot_2(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_02_cost  # noqa: F401

        assert 2 in slide_registry

    def test_registered_value_is_add_slide(self):
        from scripts.slides import slide_registry
        from scripts.slides.slide_02_cost import add_slide

        assert slide_registry[2] is add_slide

    def test_registered_value_is_callable(self):
        from scripts.slides import slide_registry
        import scripts.slides.slide_02_cost  # noqa: F401

        assert callable(slide_registry[2])


# ---------------------------------------------------------------------------
# Internal helper: _cost_per_story_by_date
# ---------------------------------------------------------------------------


class TestCostPerStoryByDate:
    def test_divides_cost_by_shipped(self):
        from scripts.slides.slide_02_cost import _cost_per_story_by_date

        rows = [{"date": "2026-04-01", "project": "TK", "shipped": 4, "cost_usd": 1.0}]
        dates, costs = _cost_per_story_by_date(rows)
        assert dates == ["2026-04-01"]
        assert abs(costs[0] - 0.25) < 1e-9

    def test_aggregates_across_projects(self):
        from scripts.slides.slide_02_cost import _cost_per_story_by_date

        rows = [
            {"date": "2026-04-01", "project": "TK", "shipped": 2, "cost_usd": 0.50},
            {"date": "2026-04-01", "project": "FA", "shipped": 2, "cost_usd": 0.30},
        ]
        dates, costs = _cost_per_story_by_date(rows)
        assert dates == ["2026-04-01"]
        # total_cost=0.80, total_shipped=4 → 0.20
        assert abs(costs[0] - 0.20) < 1e-9

    def test_shipped_zero_returns_zero_cost(self):
        from scripts.slides.slide_02_cost import _cost_per_story_by_date

        rows = [{"date": "2026-04-01", "project": "TK", "shipped": 0, "cost_usd": 0.5}]
        _, costs = _cost_per_story_by_date(rows)
        assert costs[0] == 0.0

    def test_returns_sorted_dates(self):
        from scripts.slides.slide_02_cost import _cost_per_story_by_date

        rows = [
            {"date": "2026-04-03", "project": "TK", "shipped": 1, "cost_usd": 0.1},
            {"date": "2026-04-01", "project": "TK", "shipped": 1, "cost_usd": 0.1},
            {"date": "2026-04-02", "project": "TK", "shipped": 1, "cost_usd": 0.1},
        ]
        dates, _ = _cost_per_story_by_date(rows)
        assert dates == sorted(dates)

    def test_multiple_dates_independent(self):
        from scripts.slides.slide_02_cost import _cost_per_story_by_date

        rows = [
            {"date": "2026-04-01", "project": "TK", "shipped": 2, "cost_usd": 0.60},
            {"date": "2026-04-02", "project": "TK", "shipped": 5, "cost_usd": 0.25},
        ]
        dates, costs = _cost_per_story_by_date(rows)
        assert abs(costs[0] - 0.30) < 1e-9
        assert abs(costs[1] - 0.05) < 1e-9


# ---------------------------------------------------------------------------
# Internal helper: _before_after_averages
# ---------------------------------------------------------------------------


class TestBeforeAfterAverages:
    def test_splits_on_switch_date(self):
        from scripts.slides.slide_02_cost import _before_after_averages

        dates = ["2026-04-15", "2026-04-16", "2026-04-17", "2026-04-18"]
        costs = [0.50, 0.40, 0.10, 0.05]
        before, after = _before_after_averages(dates, costs)
        assert abs(before - 0.45) < 1e-9   # (0.50 + 0.40) / 2
        assert abs(after - 0.075) < 1e-9   # (0.10 + 0.05) / 2

    def test_all_before_returns_none_after(self):
        from scripts.slides.slide_02_cost import _before_after_averages

        dates = ["2026-04-10", "2026-04-11"]
        costs = [0.50, 0.40]
        before, after = _before_after_averages(dates, costs)
        assert before is not None
        assert after is None

    def test_all_after_returns_none_before(self):
        from scripts.slides.slide_02_cost import _before_after_averages

        dates = ["2026-04-17", "2026-04-18"]
        costs = [0.10, 0.05]
        before, after = _before_after_averages(dates, costs)
        assert before is None
        assert after is not None

    def test_zero_cost_days_excluded_from_averages(self):
        from scripts.slides.slide_02_cost import _before_after_averages

        dates = ["2026-04-15", "2026-04-16"]
        costs = [0.0, 0.40]
        before, _ = _before_after_averages(dates, costs)
        assert abs(before - 0.40) < 1e-9

    def test_all_zero_before_returns_none(self):
        from scripts.slides.slide_02_cost import _before_after_averages

        dates = ["2026-04-15"]
        costs = [0.0]
        before, _ = _before_after_averages(dates, costs)
        assert before is None

    def test_switch_date_counted_as_after(self):
        from scripts.slides.slide_02_cost import _before_after_averages

        dates = ["2026-04-16", "2026-04-17"]
        costs = [0.80, 0.20]
        before, after = _before_after_averages(dates, costs)
        assert abs(before - 0.80) < 1e-9
        assert abs(after - 0.20) < 1e-9
